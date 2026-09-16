import asyncio
import tempfile
import unittest
from pathlib import Path
from unittest.mock import AsyncMock

from playwright.async_api import async_playwright
from indoteknik_scraper import (FIELDS, database, export, extract, image_extension,
                               next_page, normalize_url, save_page)


class Checkpoints(unittest.TestCase):
    def test_canonical_urls(self):
        self.assertEqual(normalize_url('https://www.indoteknik.com/shop/product/item-10?utm_source=home', product=True),
                         'https://indoteknik.com/shop/product/item-10')
        self.assertEqual(normalize_url('/shop/category/tools-1?slug=tools-1&page=1'),
                         'https://indoteknik.com/shop/category/tools-1')
        self.assertIn('page=2', normalize_url('/shop/category/tools-1?page=2'))

    def test_image_validation(self):
        self.assertEqual(image_extension(b'\x89PNG\r\n\x1a\nrest'), '.png')
        self.assertEqual(image_extension(b'RIFF1234WEBPrest'), '.webp')
        with self.assertRaises(ValueError):
            image_extension(b'<html>Error 503</html>')

    def test_resume_dedup_relationships_and_export(self):
        with tempfile.TemporaryDirectory() as folder:
            output = Path(folder)
            db = database(output)
            try:
                for url in ['category-a', 'category-b']:
                    db.execute('INSERT INTO categories(url,name,next_url) VALUES(?,?,?)', (url, url, url))
                db.commit()
                row = dict.fromkeys(FIELDS, '')
                row.update(product_id='10', product_name='Tool', product_url='product-10', price=100,
                           price_incl_ppn=111, image_status='downloaded')
                save_page(db, [row], 'category-a', 'page-2')
                save_page(db, [row], 'category-a', 'page-2')
                self.assertEqual(db.execute('SELECT done,next_url FROM categories WHERE url=?', ('category-a',)).fetchone(), (0, 'page-2'))
                save_page(db, [row], 'category-b', None)
                self.assertEqual(db.execute('SELECT COUNT(*) FROM products').fetchone()[0], 1)
                self.assertEqual(db.execute('SELECT COUNT(*) FROM relationships').fetchone()[0], 2)
                export(db, output)
                import openpyxl
                book = openpyxl.load_workbook(output / 'indoteknik_catalog.xlsx')
                self.assertEqual(book.sheetnames, ['Products', 'Brands', 'Categories', 'Product Categories'])
                self.assertEqual(book['Products'].max_row, 2)
                self.assertEqual(book['Product Categories'].max_row, 3)
                book.close()
            finally:
                db.close()


class BrowserExtraction(unittest.IsolatedAsyncioTestCase):
    async def test_prices_and_exclusion_of_banner(self):
        async with async_playwright() as p:
            browser = await p.chromium.launch()
            try:
                page = await browser.new_page()
                await page.set_content('''<a href="https://indoteknik.com/shop/product/banner-9">Promo</a>
                  <div class="rounded shadow-sm">
                    <a href="https://indoteknik.com/shop/brands/acme-1">Acme</a>
                    <a title="Tool" href="https://indoteknik.com/shop/product/tool-10?utm_source=test">Tool</a>
                    <div><s>Rp 200.000</s><span>Rp 100.000</span><div>Inc. PPN: Rp 111.000</div></div>
                    <div>Voucher: Rp 10.000</div>
                    <img class="gambarA" src="data:image/gif;base64,R0lGODlhAQABAIAAAAAAAP///yH5BAEAAAAALAAAAAABAAEAAAIBRAA7">
                  </div>''')
                rows = await extract(page)
                self.assertEqual(len(rows), 1)
                self.assertEqual(rows[0]['product_name'], 'Tool')
                self.assertEqual(rows[0]['brand'], 'Acme')
                self.assertEqual(rows[0]['price'], 100000)
                self.assertEqual(rows[0]['price_incl_ppn'], 111000)
                self.assertEqual(rows[0]['product_url'], 'https://indoteknik.com/shop/product/tool-10')
            finally:
                await browser.close()

    async def test_pagination_does_not_repeat_page_one_or_jump(self):
        page = AsyncMock()
        page.url = 'https://indoteknik.com/shop/category/tools-1'
        locator = AsyncMock()
        page.locator = lambda _: locator
        locator.evaluate_all.return_value = [
            {'text': '1', 'href': page.url + '?page=1'},
            {'text': '2', 'href': page.url + '?page=2'},
            {'text': '87', 'href': page.url + '?page=87'}]
        self.assertEqual(await next_page(page), page.url + '?page=2')
        locator.evaluate_all.return_value = [{'text': '87', 'href': page.url + '?page=87'}]
        with self.assertRaises(RuntimeError):
            await next_page(page)
        page.url += '?page=87'
        self.assertIsNone(await next_page(page))


if __name__ == '__main__':
    unittest.main()
